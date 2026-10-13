from __future__ import annotations

import argparse
import base64
import html
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


CSS = """
:root { color-scheme: light; font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", "Microsoft YaHei", sans-serif; }
body { margin: 0; background: #f4f7fa; color: #18212b; line-height: 1.65; }
main { max-width: 1100px; margin: 0 auto; padding: 20px 16px 56px; background: white; min-height: 100vh; }
h1 { color: #1f4e79; font-size: clamp(24px, 5vw, 38px); line-height: 1.25; margin: 8px 0 20px; }
h2, h3 { color: #1f4e79; line-height: 1.3; margin-top: 28px; }
p { margin: 10px 0; }
table { width: 100%; border-collapse: collapse; margin: 14px 0 20px; font-size: 14px; overflow: hidden; }
th { background: #1f4e79; color: white; text-align: left; }
th, td { border: 1px solid #cbd5df; padding: 8px 9px; vertical-align: top; }
tr:nth-child(even) { background: #f1f6fa; }
img { display: block; max-width: 100%; height: auto; margin: 14px auto; }
.meta { color: #687786; font-size: 13px; border-bottom: 1px solid #d9e1e8; padding-bottom: 12px; }
@media (max-width: 680px) { main { padding: 14px 10px 40px; } table { display: block; overflow-x: auto; white-space: normal; } th, td { min-width: 110px; } p { font-size: 15px; } }
"""


def image_data(doc: Document, rid: str) -> str | None:
    part = doc.part.related_parts.get(rid)
    if not part or not hasattr(part, "blob"):
        return None
    content_type = getattr(part, "content_type", "image/png")
    encoded = base64.b64encode(part.blob).decode("ascii")
    return f"data:{content_type};base64,{encoded}"


def paragraph_html(doc: Document, paragraph) -> str:
    style = paragraph.style.name if paragraph.style else ""
    text = html.escape(paragraph.text.strip())
    if not text and not paragraph._p.xpath(".//a:blip"):
        return ""
    for blip in paragraph._p.xpath(".//a:blip"):
        rid = blip.get(qn("r:embed"))
        src = image_data(doc, rid) if rid else None
        if src:
            text += f'<img src="{src}" alt="报告图表">'
    if style.startswith("Heading 1"):
        return f"<h1>{text}</h1>"
    if style.startswith("Heading 2"):
        return f"<h2>{text}</h2>"
    if style.startswith("Heading"):
        return f"<h3>{text}</h3>"
    return f"<p>{text}</p>"


def table_html(table) -> str:
    rows = []
    for row_index, row in enumerate(table.rows):
        cells = []
        tag = "th" if row_index == 0 else "td"
        for cell in row.cells:
            value = "<br>".join(html.escape(p.text.strip()) for p in cell.paragraphs if p.text.strip())
            cells.append(f"<{tag}>{value}</{tag}>")
        rows.append("<tr>" + "".join(cells) + "</tr>")
    return "<table>" + "".join(rows) + "</table>"


def convert(source: Path, target: Path) -> None:
    doc = Document(source)
    blocks = []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = next((p for p in doc.paragraphs if p._p is child), None)
            if paragraph is not None:
                blocks.append(paragraph_html(doc, paragraph))
        elif child.tag.endswith("}tbl"):
            table = next((t for t in doc.tables if t._tbl is child), None)
            if table is not None:
                blocks.append(table_html(table))
    title = html.escape(doc.core_properties.title or source.stem)
    document = "<!doctype html><html lang=\"zh-CN\"><head><meta charset=\"utf-8\"><meta name=\"viewport\" content=\"width=device-width, initial-scale=1\"><title>" + title + "</title><style>" + CSS + "</style></head><body><main>" + "".join(blocks) + "<p class=\"meta\">来源文件：" + html.escape(source.name) + "</p></main></body></html>"
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(document, encoding="utf-8")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("target", type=Path)
    args = parser.parse_args()
    convert(args.source, args.target)
